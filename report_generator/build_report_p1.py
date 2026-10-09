"""
Script to generate the complete, exhaustive Academic & Technical Project Report for:
Wedding Planner and Vendor Coordination Platform (WPVCP)
Vows & Veils — Enterprise Wedding Concierge & Coordination Suite

Course: 23CS54C – MODERN WEB TECHNOLOGIES
College: NATIONAL ENGINEERING COLLEGE, KOVILPATTI
Student: HARIRAAM B (24104085)
"""

import os
import json
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

FIGURES_DIR = r"c:\Users\Welcome\Downloads\WPVCP Modified\report_generator\figures"
METADATA_FILE = r"c:\Users\Welcome\Downloads\WPVCP Modified\report_generator\schema_metadata.json"
OUTPUT_PATH = r"c:\Users\Welcome\Downloads\Wedding_Planner_and_Vendor_Coordination_Platform_Report.docx"
WORKSPACE_OUTPUT = r"c:\Users\Welcome\Downloads\WPVCP Modified\Wedding_Planner_and_Vendor_Coordination_Platform_Report.docx"

# Helper XML functions
def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def make_header_repeat(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = parse_xml(f'<w:tblHeader {nsdecls("w")}/>')
    trPr.append(tblHeader)

def make_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    cantSplit = parse_xml(f'<w:cantSplit {nsdecls("w")}/>')
    trPr.append(cantSplit)

def add_styled_table(doc, headers, data, col_widths=None, align_cols=None):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    # Style Header Row
    hdr_cells = table.rows[0].cells
    make_header_repeat(table.rows[0])
    make_row_cant_split(table.rows[0])
    
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        set_cell_background(hdr_cells[i], "1E293B")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if align_cols and align_cols[i] == 'C' else WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

    # Style Data Rows
    for r_idx, row_data in enumerate(data):
        row = table.rows[r_idx + 1]
        make_row_cant_split(row)
        bg_col = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.text = str(val)
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, top=90, bottom=90, left=140, right=140)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if align_cols and align_cols[c_idx] == 'C' else WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(9.0)

    # Set column widths
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)

    # Spacing after table
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(4)
    p_after.paragraph_format.space_after = Pt(6)
    return table

def add_p(doc, text="", bold=False, italic=False, size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6, font_name="Times New Roman", color=None):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = font_name
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
    return p

def add_bullet(doc, title="", desc="", size=12):
    p = doc.add_paragraph(style='List Paragraph')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if title:
        r1 = p.add_run(title + ": ")
        r1.bold = True
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(size)
    if desc:
        r2 = p.add_run(desc)
        r2.bold = False
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(size)
    return p

def add_ch_title(doc, ch_num, ch_title):
    doc.add_page_break()
    p_ch = doc.add_paragraph()
    p_ch.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ch.paragraph_format.space_before = Pt(18)
    p_ch.paragraph_format.space_after = Pt(4)
    r_ch = p_ch.add_run(f"CHAPTER {ch_num}")
    r_ch.bold = True
    r_ch.font.name = "Times New Roman"
    r_ch.font.size = Pt(16)

    p_t = doc.add_paragraph()
    p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t.paragraph_format.space_before = Pt(2)
    p_t.paragraph_format.space_after = Pt(16)
    r_t = p_t.add_run(ch_title.upper())
    r_t.bold = True
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(16)

def add_sec_h2(doc, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    return p

def add_sec_h3(doc, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    return p

def add_figure(doc, img_name, caption_text, width_inches=6.0):
    img_path = os.path.join(FIGURES_DIR, img_name)
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        run = p_img.add_run()
        run.add_picture(img_path, width=Inches(width_inches))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run(caption_text)
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(10)
        r_cap.font.bold = True
        r_cap.font.color.rgb = RGBColor(71, 85, 105)
    else:
        print(f"Warning: Image file not found: {img_path}")

print("Helper functions initialized.")
